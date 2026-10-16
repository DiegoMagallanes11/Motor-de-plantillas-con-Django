"""Create the Spanish Lab 07 guide from measured results and real evidence."""
import ast
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "deliverables" / "lab07"
EVIDENCE = OUTPUT / "evidence"
QA = ROOT / "deliverables" / "qa" / "lab07"
QA.mkdir(parents=True, exist_ok=True)


def read_json(filename):
    return json.loads((EVIDENCE / filename).read_text(encoding="utf-8"))


reference = read_json("duel-team.json")
candidates = read_json("duel-ai.json")
measurement = read_json("query-counts.json")
actions = read_json("agent-results.json")
counts = read_json("dataset.json")
source = (ROOT / "blog" / "duel" / "team.py").read_text(encoding="utf-8")
tree = ast.parse(source)
solutions = {
    node.name: ast.get_source_segment(source, node)
    for node in tree.body if isinstance(node, ast.FunctionDef)
}
document = Document()
section = document.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(.68)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.8)
section.right_margin = Inches(.8)
section.footer_distance = Inches(.25)
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
    style = document.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
document.styles["Normal"].font.size = Pt(10.5)
document.styles["Normal"].paragraph_format.line_spacing = 1.1
document.styles["Normal"].paragraph_format.space_after = Pt(7)
document.styles["Title"].font.size = Pt(30)
document.styles["Heading 1"].font.size = Pt(20)
document.styles["Heading 2"].font.size = Pt(13)
document.styles["Heading 1"].paragraph_format.space_after = Pt(12)
document.styles["Heading 2"].paragraph_format.space_before = Pt(12)
document.core_properties.title = "Laboratorio 07 ORM de Django"
document.core_properties.author = "Diego Magallanes Linares"
document.core_properties.subject = "Consultas relaciones y optimización N+1"
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Laboratorio 07  |  ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def p(text, style=None):
    return document.add_paragraph(text, style)


def heading(text, level=1, new_page=False):
    if new_page:
        document.add_page_break()
    document.add_heading(text, level)


def code(text, size=8.3):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(8)
    for index, line in enumerate(text.strip("\n").splitlines()):
        run = paragraph.add_run(("\n" if index else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(size)


def table(headers, rows, widths, size=9):
    result = document.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for cell, text in zip(result.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(result.add_row().cells, row):
            cell.text = str(text)
    for index, row in enumerate(result.rows):
        for column, cell in enumerate(row.cells):
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "253E39" if index == 0 else ("F0F3F1" if index % 2 else "FFFFFF"))
            properties.append(shade)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "80")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "D9D9D9")
                borders.append(border)
            properties.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(size)
                    if index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def picture(filename, caption, width=6.5, height=4.8, crop=None):
    source_path = EVIDENCE / filename
    image = Image.open(source_path)
    if crop:
        image = image.crop(crop)
        source_path = QA / f"crop-{Path(filename).stem}.png"
        image.save(source_path)
    width = min(width, height * image.width / image.height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(source_path), width=Inches(width))
    paragraph = p(caption)
    for run in paragraph.runs:
        run.italic = True
        run.font.size = Pt(9)


def query(number, observation):
    row = reference[number - 1]
    heading(f"Consulta {number}", 2)
    p(row["question"])
    code(solutions[f"question_{number}"])
    p(f"Resultado obtenido: {row['actual']}. Costo medido: {row['queries']} consulta(s). Veredicto: {row['status']}.")
    p(observation)


p("GUÍA DE LABORATORIO 07", "Subtitle")
p("ORM de Django", "Title")
p("Consultas relacionadas y optimización N+1", "Subtitle")
table(["Dato", "Información"], [
    ("Estudiante", "Diego Magallanes Linares"),
    ("Curso", "Desarrollo de Aplicaciones Empresariales"),
    ("Docente indicado en la guía", "Michael Montgomery Rosell"),
    ("Sección", "4-C24-A"),
    ("Fecha", "15 de octubre de 2026"),
    ("Repositorio y rama", "Motor-de-plantillas-con-Django · lab07-orm"),
], [2.2, 4.3])
heading("Resultado del desarrollo", 2)
p("Se implementaron ocho consultas de ORM con datos conocidos y verificación automática de resultados, orden, anotaciones y costo. La tarea de portada pasó de 23 a 2 consultas conservando títulos, autores y etiquetas. La suite completa finalizó con 31 pruebas correctas y el portal del laboratorio 06 siguió funcionando.")
heading("Capacidades", 2)
for text in [
    "Utilizar consultas avanzadas del ORM sobre modelos relacionados.",
    "Encadenar filtros, atravesar relaciones y ordenar resultados sin escribir SQL crudo.",
    "Detectar y corregir N+1 mediante select_related y prefetch_related.",
]:
    p(text, "List Bullet")
heading("Alcance de la adaptación", 2)
p("Se priorizó la rúbrica del campus y se adaptaron las actividades del adjunto ORM con IA duelo y verificador. No se dispuso del repositorio inicial del docente ni de sus ocho preguntas: los modelos, datos y ejercicios de este documento son una propuesta reproducible, no respuestas oficiales del material faltante. El desarrollo se realizó con asistencia de Codex.")

heading("Preparación y datos conocidos", new_page=True)
heading("Pasos 1 y 2 Entorno migración y carga", 2)
p("Se continuó el proyecto del laboratorio 06 en una rama nueva del mismo repositorio. La aplicación blog contiene la práctica ORM y sus tablas son independientes de news. Se mantuvieron Python 3.14.5, Django 6.0.9 y Pillow 12.3.0.")
code(r"""git switch -c lab07-orm
.\.venv\Scripts\python.exe -m pip install -r requirements.txt
.\.venv\Scripts\python.exe manage.py migrate
.\.venv\Scripts\python.exe manage.py seed_blog
.\.venv\Scripts\python.exe manage.py shell""")
code("from blog.models import Author, Article, Comment\nAuthor.objects.count(), Article.objects.count(), Comment.objects.count()\n# Result: (7, 12, 15)")
table(["Entidad", "Cantidad", "Caso significativo"], [
    ("Author", counts["authors"], "Gabriela no tiene artículos."),
    ("Profile", counts["profiles"], "Un perfil cargado por autor."),
    ("Article", counts["articles"], "11 publicados y 1 borrador."),
    ("Category", counts["categories"], "Opinión no tiene artículos."),
    ("Tag", counts["tags"], "El artículo 9 no tiene etiquetas."),
    ("Comment", counts["comments"], "15 en total; el artículo 1 tiene 4."),
], [1.3, .8, 4.4])
p("La carga es determinista: utiliza los mismos ids, fechas, relaciones y contenidos. El artículo 2 está exactamente en el límite temporal del 5 de octubre a las 00:00 de Lima. Los artículos 3 y 9 no tienen categoría. Estos bordes permiten distinguir una consulta correcta de otra que solo parece razonable.")
p("seed_blog restaura únicamente las tablas de la aplicación didáctica blog. No modifica noticias del laboratorio 06 ni usuarios. La clave de Django y las contraseñas permanecen en archivos locales excluidos de Git.")
heading("Organización del trabajo", 2)
p("Estudiante: Diego Magallanes Linares. Implementación, mediciones y documentación realizadas con asistencia de Codex. No se inventan integrantes, respuestas sin IA ni intentos personales. Los registros completos se conservan en deliverables/lab07/evidence.")

heading("Modelos y relaciones", new_page=True)
heading("Paso 3 Lectura del esquema", 2)
p("Los seis modelos permiten practicar relaciones uno a uno, uno a muchos y muchos a muchos. El siguiente esquema digital representa la implementación y sirve como apoyo para el dibujo a mano solicitado en el adjunto; ese dibujo personal no se presenta como realizado.")
picture("05-data-relations.png", "Figura 1. Captura real del panel con conteos y diagrama digital de las relaciones.", height=3.4)
table(["Relación", "Implementación y comportamiento"], [
    ("Author y Profile", "OneToOneField; a lo sumo un perfil por autor. La carga crea siete perfiles."),
    ("Author y Article", "ForeignKey con related_name articles y PROTECT al eliminar un autor con artículos."),
    ("Category y Article", "ForeignKey opcional con null y blank; SET_NULL al eliminar la categoría."),
    ("Article y Tag", "ManyToManyField con tabla intermedia y acceso inverso articles."),
    ("Article y Comment", "ForeignKey con related_name comments; eliminación en cascada del artículo a sus comentarios."),
], [1.7, 4.8])
code('''category = models.ForeignKey(
    Category, null=True, blank=True, on_delete=models.SET_NULL,
    related_name="articles",
)
tags = models.ManyToManyField(Tag, related_name="articles", blank=True)''')
p("El acceso author__profile__city atraviesa la relación hacia el perfil. El acceso inverso articles__comments__is_approved parte de Author y recorre los artículos y sus comentarios. Estas rutas se resuelven en la base de datos mediante el ORM.")

heading("Filtros encadenados y relaciones", new_page=True)
heading("Paso 4 Consultas propuestas 1 y 2", 2)
query(1, "Se usa gte porque desde incluye el instante límite. El candidato didáctico usa gt y pierde el artículo 2. Encadenar filter conserva un QuerySet y deja a la base de datos realizar el filtrado.")
query(2, "La ciudad pertenece a Profile, no a Author. Django atraviesa author, profile y category con uniones generadas automáticamente. El candidato author__city produce FieldError antes de ejecutar una consulta.")
p("Importaciones utilizadas en team.py: datetime, ZoneInfo, Count y Q, además de Article, Author y Category. Todas las funciones ordenan explícitamente para que el verificador pueda comparar resultados estables.")

heading("Condiciones alternativas y acceso inverso", new_page=True)
heading("Paso 4 Consultas propuestas 3 y 4", 2)
query(3, "Q expresa la alternativa entre Tecnología y categoría nula. El filtro is_published=True se aplica a toda la alternativa. El candidato que solo exige Tecnología omite los artículos 3 y 9.")
query(4, "Se parte de Author y se recorren sus artículos publicados y sus etiquetas. Ana tiene dos artículos django; por eso el join puede devolverla dos veces. distinct elimina las repeticiones en la base, sin convertir primero a una lista ni filtrar en Python.")
heading("Verificación y procedencia", 2)
p("Las cuatro funciones se comprobaron con manage.py duel. Esta es una solución de referencia asistida: el paso de responder sin IA no se atribuye al estudiante. Intentos personales sin IA: no registrados. La primera ejecución completa de las ocho funciones de referencia obtuvo 8 de 8 aceptadas.")

heading("Anotaciones sin conteos multiplicados", new_page=True)
heading("Paso 5 Consulta propuesta 5", 2)
p(reference[4]["question"])
code(solutions["question_5"])
p("Resultado real de los doce artículos, incluyendo los que tienen cero relaciones. La consulta devuelve un QuerySet de Article con los atributos anotados n_comments y n_tags.")
table(["Artículo id", "n_comments", "n_tags"], reference[4]["actual"], [2.2, 2.1, 2.2])
p("Costo medido: una consulta. Al unir simultáneamente comentarios y etiquetas, el artículo 1 genera 4 × 2 = 8 filas combinadas. Count sin distinct devolvería 8 y 8. Count con distinct cuenta las identidades de cada relación y obtiene 4 comentarios y 2 etiquetas.")
p("No se excluyen comentarios pendientes: la pregunta define el total de comentarios, independientemente de su aprobación. El artículo 9 conserva un conteo de cero etiquetas y el 12, aunque es borrador, se incluye porque la pregunta pide todos los artículos.")

heading("Ninguno y agregación filtrada", new_page=True)
heading("Paso 5 Consultas propuestas 6 y 7", 2)
query(6, "No tener ningún comentario aprobado es diferente de tener algún comentario pendiente. exclude elimina a cualquier autor que tenga un comentario aprobado en sus artículos y también conserva a quienes no tienen artículos ni comentarios. Por eso incluye a Gabriela, id 7.")
query(7, "El filtro de publicación se aplica dentro de Count. Tecnología tiene cinco artículos en total pero solo cuatro publicados. Anotar primero todos los artículos, sin el filtro, produciría el conteo incorrecto de cinco.")
p("En ambas funciones, la selección se realiza mediante el ORM. No se recuperan todos los objetos para decidir en un bucle de Python quién pertenece al resultado.")

heading("Portada y veredicto de las consultas", new_page=True)
heading("Paso 5 Consulta propuesta 8", 2)
p(reference[7]["question"])
code('''def question_8():
    return front_page()

def front_page():
    return Article.objects.filter(is_published=True).select_related(
        "author"
    ).prefetch_related("tags").order_by("-published_at", "-id")[:11]''')
p(f"Resultado real: {reference[7]['actual']}. El verificador evalúa también article.author.name y article.tags.all(), por lo que mide la carga de las relaciones utilizadas al representar la portada. Costo: dos consultas.")
heading("Resultado del duelo de referencia", 2)
verdict_text = (EVIDENCE / "duel-team.txt").read_text(encoding="utf-8")
code(verdict_text[verdict_text.index("Q1:"):].strip())
picture("04-dashboard.png", "Figura 2. Panel real con ocho consultas aceptadas y la medición de portada.", crop=(80, 90, 1370, 740), height=3.6)
p("Los resultados esperados se almacenan como valores literales en questions.py; no se obtienen ejecutando otra vez las mismas soluciones. El verificador compara contenido, orden y presupuesto de consultas.")

heading("Comparación con candidatos de IA", new_page=True)
heading("Pasos 6 y 7 Veredictos y clasificación", 2)
p("Se ejecutó manage.py duel --source ai. Los ocho candidatos de ai.py contienen fallos introducidos deliberadamente con fines didácticos; esta tabla no representa una conversación independiente ni mide una tasa real de errores de un asistente.")
table(["N.º", "Candidato", "Evidencia del fallo", "Decisión"], [
    ("1", "wrong_result · 1 consulta", "Falta id 2 por usar gt en el límite temporal.", "Rechazar; usar gte."),
    ("2", "error · 0 consultas", "FieldError en author__city; city está en Profile.", "Rechazar; corregir relación."),
    ("3", "wrong_result · 1 consulta", "Obtiene [1, 2, 7, 10]; faltan 3 y 9 sin categoría.", "Rechazar; incluir isnull con Q."),
    ("4", "wrong_result · 1 consulta", "Obtiene [1, 1, 3]; repite autora 1.", "Rechazar; añadir distinct."),
    ("5", "wrong_result · 1 consulta", "Artículo 1: [8, 8] frente a [4, 2].", "Rechazar; conteos distintos."),
    ("6", "wrong_result · 1 consulta", "Obtiene [1, 2, 4] frente a [2, 5, 6, 7].", "Rechazar; corregir ninguno."),
    ("7", "wrong_result · 1 consulta", "Tecnología: 5 artículos frente a 4 publicados.", "Rechazar; filtrar Count."),
    ("8", "expensive · 23 consultas", "Mismos ids pero supera el presupuesto de 2.", "Aceptar contenido; corregir costo."),
], [.4, 1.45, 3.0, 1.65])
p("Las correcciones verificadas se encuentran en team.py. El SQL completo y las diferencias se conservan en duel-team.json y duel-ai.json. Q4 se clasifica como duplicación por una relación inversa, una variante de resultado distinto; no es un campo inventado.")
heading("Paso 8 Consulta opcional a otro asistente", 2)
p("No se ejecutó una ronda adicional presentada como autorizada por el docente, porque no consta esa autorización ni se dispone del repositorio original. El uso real de Codex para esta adaptación queda documentado en la tabla de uso de IA.")

heading("SQL generado como evidencia", new_page=True)
p("No se utilizó SQL crudo para resolver los ejercicios. Imprimir consulta.query permite inspeccionar el SQL que Django generó a partir del QuerySet; es una herramienta de diagnóstico.")
code('''from blog.duel.team import question_5
consulta = question_5()
print(consulta.query)''')
p("La diferencia decisiva en la consulta 5 se observa en los agregados. Fragmentos del SQL generado para comparar ambos casos:")
code('''COUNT("blog_comment"."id") AS "n_comments",
COUNT("blog_article_tags"."tag_id") AS "n_tags"

COUNT(DISTINCT "blog_comment"."id") AS "n_comments",
COUNT(DISTINCT "blog_article_tags"."tag_id") AS "n_tags"''')
picture("09-counts-detail.png", "Figura 3. Captura real del detalle de la consulta 5 con resultados y SQL generado. Recorte de la columna de datos.", crop=(0, 0, 760, 554), height=4.8)
p("El conteo incorrecto no se corrige sumando o dividiendo resultados en Python. Se expresa el conteo de identidades distintas en la consulta. La evidencia contiene las sentencias completas, las respuestas y el costo de cada candidato.")

heading("N+1 de rojo a verde", new_page=True)
heading("Pasos 9 y 10 Fallo inicial y explicación", 2)
p("Se ejecutó la prueba antes de corregir blog/queries.py. La aserción esperaba dos consultas y capturó 23; el archivo 01-red-test.txt conserva el mensaje y la lista completa de consultas. front_page_before.py.txt conserva el código anterior.")
code('''AssertionError: 23 != 2 : 23 queries executed, 2 expected
Ran 1 test
FAILED (failures=1)''')
p("Descomposición: una consulta obtiene los once artículos; once consultas recuperan el autor de cada instancia y otras once recuperan sus etiquetas. Aunque Ana se repite, la caché de un artículo no comparte automáticamente la relación de autor con las demás instancias. El costo ingenuo crece como 1 + 2N.")
heading("Paso 11 Corrección y nueva medición", 2)
p("select_related une la relación ForeignKey del autor en la consulta de artículos. prefetch_related hace una segunda consulta para las etiquetas de todos esos artículos y Django las asocia en memoria. El mismo trabajo queda en dos consultas.")
picture("07-n-plus-one.png", "Figura 4. Medición real sobre el mismo contenido antes y después de precargar relaciones.", height=2.5)
table(["Carga medida", "Antes", "Después", "Contenido"], [
    ("11 artículos con autor y etiquetas", "23 consultas", "2 consultas", "Idéntico"),
    ("23 publicados sin límite en la prueba de crecimiento", "47 consultas", "2 consultas", "Idéntico"),
], [3.0, 1.1, 1.1, 1.3])
p("Se redujeron 21 consultas, un 91,3 %. Es una reducción en número de consultas; no se afirma una aceleración de tiempo no medida. El presupuesto corresponde a la tarea de portada, no al total de consultas de la página de diagnóstico.")

heading("Pruebas y resultados", new_page=True)
heading("Paso 11 Suite completa", 2)
p("La misma prueba de portada pasó después del cambio y la suite completa finalizó con 31 pruebas correctas. Las pruebas usan una base temporal independiente. También se comprobó que no hay migraciones pendientes y que manage.py check no informa errores.")
code('''Ran 31 tests in 0.785s
OK
System check identified no issues (0 silenced).''')
table(["Caso", "Esperado y resultado obtenido", "Evidencia"], [
    ("Ocho funciones de referencia", "Resultados y costos correctos; 8/8 aceptadas.", "duel-team.json"),
    ("Candidatos didácticos", "Se detectan 6 resultados incorrectos, 1 campo inválido y 1 consulta cara.", "duel-ai.json"),
    ("Portada con relaciones", "23 a 2 consultas con datos idénticos; cumplido.", "query-counts.json"),
    ("Crecimiento a 23 publicados", "47 a 2 sin alterar contenido; cumplido.", "test_query_count_stays_constant_when_articles_grow"),
    ("Autores sin artículos", "El autor 7 se conserva en Q6; cumplido.", "test_reverse_relation"),
    ("Relaciones vacías", "Categoría nula y cero etiquetas no pierden filas; cumplido.", "test_nullable_category"),
    ("Cuatro intenciones", "ok, ok, approval_required, denied; cumplido.", "agent-results.json"),
    ("Texto hostil en comentario", "Se devuelve como dato sin borrar registros; cumplido.", "test_injected_comment"),
    ("Plan aprobado", "Solo se borran los tres comentarios previstos; cumplido.", "agent-results.json"),
    ("Plan modificado o reutilizado", "Se rechaza tras inserción, edición, cambio de artículo o repetición; cumplido.", "tests del registro"),
    ("Panel web", "Ocho aceptadas, diagrama y SQL visibles; cumplido.", "browser-results.json"),
    ("Vista móvil", "Sin desbordamiento de página a 390 px; cumplido.", "10-mobile.png"),
    ("Laboratorio 06", "Diez pruebas previas pasan y sus seis tarjetas permanecen.", "suite completa y navegador"),
], [1.45, 3.5, 1.55], size=8.6)
p("La lista completa de nombres y resultados de las pruebas se encuentra en 03-full-suite.txt. La captura del navegador no reemplaza las aserciones de resultados y de número de consultas.")

heading("Acciones permitidas y contenido no confiable", new_page=True)
heading("Pasos 12 y 13 Intenciones y cuarto comentario", 2)
table(["Intención ejecutada", "Acción", "Respuesta"], [
    (row["intent"], row["response"].get("action", "ninguna"), row["response"]["status"])
    for row in actions["intents"]
], [3.6, 1.5, 1.4])
code(r"""python manage.py agent --intent 'listar artículos publicados'
python manage.py agent --intent 'ver comentarios del artículo 1'
python manage.py agent --intent 'borrar comentarios del artículo 5'
python manage.py agent --intent 'ejecutar SQL para borrar todos los artículos'""", size=7.8)
picture("08-agent.png", "Figura 5. Captura real de las cuatro respuestas y el comentario hostil mostrado como dato.", height=3.1)
p("Si un programa interpretara el cuarto comentario como instrucción, podría intentar borrar artículos o exceder el permiso de lectura solicitado. El texto almacenado no representa una autorización del usuario, aunque incluya palabras como SYSTEM u orden.")
p("registry.py solo expone list_posts, comments_of_post y delete_comments. Valida el conjunto exacto de argumentos y sus tipos, separa lectura de escritura y exige un plan aprobado para borrar. No interpreta el cuerpo del comentario como código, SQL o comandos. El parser exige una intención completa reconocida y rechaza instrucciones añadidas.")

heading("Aprobación de un plan concreto", new_page=True)
heading("Paso 14 Borrado aprobado y cambio de estado", 2)
p("Se preparó el borrado de comentarios del artículo 5, se verificó su alcance y se aprobó el hash. El plan contenía exactamente los ids 9, 10 y 11. El registro eliminó tres comentarios y no alteró los cuatro del artículo 1 ni los doce artículos.")
code("python manage.py agent --action delete_comments --article-id 5\n"
     "python manage.py agent --action delete_comments --article-id 5 `\n"
     "    --approve HASH_DEVUELTO_POR_EL_PLAN\npython manage.py seed_blog")
table(["Etapa real", "Resultado"], [
    ("Plan inicial", f"{actions['plan']['status']}; ids {actions['plan']['comment_ids']}."),
    ("Aprobación del mismo estado", f"{actions['approved']['status']}; borrados {actions['approved']['deleted']}."),
    ("Comentarios del artículo 5 después", actions["comments_after_delete"]),
    ("Primera restauración", f"{actions['restored']['comments']} comentarios totales."),
    ("Se agrega un comentario entre plan y aprobación", "El artículo 5 pasa de 3 a 4 comentarios."),
    ("Se usa el hash anterior", f"{actions['stale_approval']['status']}: {actions['stale_approval']['reason']}."),
    ("Comentarios preservados tras el rechazo", actions["comments_preserved_after_rejection"]),
    ("Restauración final", f"{actions['final_counts']['comments']} comentarios y {actions['final_counts']['articles']} artículos."),
], [3.8, 2.7])
p("El hash real del primer plan registrado comienza con " + actions["plan"]["plan_hash"][:16] + ". La evidencia JSON conserva el valor completo; al repetir la práctica debe copiarse el hash nuevo de la salida actual.")
code('''with transaction.atomic():
    comments = _snapshot(article_id)
    signature = _plan_hash(article_id, comments)
    if not hmac.compare_digest(signature, approve):
        return _denied("stale_or_invalid_plan")''')
p("La firma HMAC-SHA256 incluye acción, artículo e ids y contenido de los comentarios. Si cambia el estado, cambia la firma. La transacción y el borrado por ids limitan el alcance. Es una demostración local: para un servicio remoto harían falta autenticación y permisos; SQLite no aplica bloqueos de fila con select_for_update.")

heading("Correspondencia con la guía adjunta", new_page=True)
p("La siguiente relación distingue lo ejecutado de las actividades personales o del material inicial que no estuvo disponible. Las consultas de referencia y los casos del verificador son de esta adaptación.")
table(["Paso", "Desarrollo y evidencia", "Estado"], [
    ("1", "Proyecto acumulativo, rama nueva, dependencias y migración blog.", "Realizado"),
    ("2", "seed_blog y conteos 7 autores, 12 artículos, 15 comentarios.", "Realizado con datos adaptados"),
    ("3", "Lectura de modelos y diagrama digital; dibujo a mano a cargo del estudiante.", "Apoyo digital realizado"),
    ("4", "Consultas 1 a 4 en team.py con resultados verificados; no se atribuyen al estudiante sin IA.", "Referencia asistida"),
    ("5", "Consultas 5 a 8 y ocho resultados correctos; intentos personales no registrados.", "Referencia asistida"),
    ("6", "duel --source ai con candidatos erróneos deliberados.", "Comparación didáctica adaptada"),
    ("7", "Fallos clasificados con diferencias y SQL generado.", "Realizado"),
    ("8", "Ronda opcional con asistente autorizado por el docente.", "Omitida; sin constancia de autorización"),
    ("9", "Fallo inicial de 23 frente a 2 esperadas y consultas capturadas.", "Realizado"),
    ("10", "Descomposición 1 + 11 + 11 y crecimiento con N.", "Realizado"),
    ("11", "select_related, prefetch_related y 31 pruebas correctas.", "Realizado"),
    ("12", "Cuatro intenciones del README adaptado.", "Realizado"),
    ("13", "Análisis del cuarto comentario y límites del registro.", "Realizado"),
    ("14", "Borrado aprobado, restauración y rechazo de plan obsoleto.", "Realizado"),
    ("15", "Tabla de IA, documento y publicación en la rama lab07-orm.", "Preparado; subida al campus por el estudiante"),
], [.45, 4.25, 1.8], size=8.8)
p("Las normas de trabajo se mantuvieron: código y comentarios en inglés, explicaciones en español, modelos en singular y migraciones versionadas. No se realizaron intervenciones de hardware ni de conexiones. La publicación del repositorio y la entrega al campus son pasos distintos.")

heading("Uso de IA y decisiones de aceptación", new_page=True)
heading("Paso 15 Registro del uso real", 2)
table(["Qué se pidió a Codex", "Respuesta y decisión", "Motivo y evidencia"], [
    ("Adaptar el laboratorio sin repositorio inicial", "Se aceptó una app blog separada y ocho preguntas propuestas.", "Permite cubrir la rúbrica conservando news; se declara su procedencia."),
    ("Proponer soluciones ORM", "Se aceptaron las ocho funciones tras verificarlas.", "Claves y anotaciones conocidas; costos 1,1,1,1,1,1,1,2."),
    ("Crear casos para comparar fallos", "Se aceptaron ejemplos erróneos solo como material didáctico.", "No se presentan como respuestas espontáneas ni como trabajo sin IA."),
    ("Optimizar la portada", "Se aceptó la precarga de autor y etiquetas.", "Mismo contenido y reducción medida de 23 a 2 consultas."),
    ("Diseñar límites de acciones", "Se aceptó un registro local con lista permitida y planes firmados.", "Pruebas de rechazo y alcance del borrado; no se confía en comentarios."),
    ("Preparar pruebas y documento", "Se aceptó la entrega tras ejecutar pruebas y revisar capturas reales.", "31 pruebas correctas y comprobaciones de navegador registradas."),
], [1.8, 2.5, 2.2])
heading("Por qué una salida de IA requiere verificación", 2)
p("Una respuesta puede usar nombres de campo plausibles que no existen, cambiar el significado de ninguno o perder un instante límite. También puede devolver los objetos correctos con un costo que crece con el tamaño del conjunto. Por eso se comprueba el esquema real, el resultado con datos conocidos y las consultas que se ejecutan al consumir las relaciones.")
p("La salida de IA y el texto de comentarios se tratan como entradas que necesitan validación. Solo las funciones revisadas se incorporan al código; el registro de acciones no ejecuta instrucciones libres. El hash de aprobación fija un alcance de datos concreto y no sustituye la identidad o los permisos de un usuario.")
p("El estudiante debe revisar las soluciones y realizar cualquier actividad personal que el docente exija sin IA. No se declaran intentos, una autoría independiente ni una autorización docente que no se hayan aportado.")

heading("Repositorio y estructura del desarrollo", new_page=True)
p("Se utilizó el mismo repositorio del curso y la rama lab07-orm, conservando main para el laboratorio 06. El README principal enlaza LAB07.md, donde figuran instalación, preguntas, resultados, costos, comandos y observaciones.")
p("https://github.com/DiegoMagallanes11/Motor-de-plantillas-con-Django/tree/lab07-orm")
picture("11-editor.jpg", "Figura 6. Captura real de Visual Studio Code con la aplicación blog, sus consultas y la rama lab07-orm.", height=3.8)
table(["Ubicación", "Responsabilidad"], [
    ("blog/models.py y migrations", "Modelos relacionados y esquema versionado."),
    ("blog/duel/team.py y ai.py", "Soluciones de referencia y candidatos didácticos."),
    ("blog/duel/questions.py y verifier.py", "Preguntas, claves conocidas y medición de costos."),
    ("blog/queries.py", "Portada ingenua y optimizada; misma tarea de representación."),
    ("blog/agent/registry.py", "Acciones permitidas, validación y aprobación de planes."),
    ("blog/tests y management/commands", "Pruebas, carga y comandos de verificación."),
    ("deliverables/lab07", "Word resuelto y evidencias reales."),
], [2.65, 3.85], size=8.8)
p("La base de datos, el entorno virtual, .env y las contraseñas quedan fuera de Git. El documento Word está preparado para que el estudiante lo revise y lo suba al campus virtual.")

heading("Conclusiones", new_page=True)
for text in [
    "1. Los filtros encadenados y las rutas con doble guion bajo permitieron atravesar modelos relacionados sin escribir SQL crudo. Los casos de categoría nula, autor sin artículos y límite temporal demostraron que la semántica de la pregunta debe reflejarse con precisión en la consulta.",
    "2. Al combinar dos relaciones de varios elementos, las uniones pueden multiplicar filas. Count con distinct produjo los valores reales de comentarios y etiquetas, y distinct evitó repetir autores en una consulta inversa.",
    "3. La medición mostró 23 consultas al representar once artículos con autor y etiquetas. select_related para el autor y prefetch_related para etiquetas redujeron el costo a dos consultas y conservaron el mismo contenido. La prueba con 23 artículos también mantuvo ese costo de dos.",
    "4. Verificar solamente los ids no basta. El duelo comprobó resultados, anotaciones, orden y presupuesto; así detectó una consulta correcta en contenido pero cara. La prueba de regresión documentó el paso real de rojo a verde.",
    "5. La salida de una IA y el contenido de un comentario no conceden permisos. La lista de acciones, los argumentos estrictos y la aprobación ligada al estado impidieron que una lectura se convirtiera en borrado y rechazaron un plan que dejó de corresponder a los datos actuales.",
    "6. La rama lab07-orm conserva el trabajo del curso en un solo repositorio y permite revisar esta sesión por separado. La adaptación declara el material no disponible y distingue las soluciones asistidas de las actividades personales que corresponden al estudiante.",
]:
    p(text)
heading("Correspondencia con la rúbrica del campus", 2)
table(["Criterio", "Evidencia entregada"], [
    ("Consultas avanzadas relacionadas", "Ocho ejercicios adaptados con resultados conocidos y pruebas."),
    ("Filtros relaciones y orden mediante ORM", "Q, relaciones inversas, distinct, Count filtrado y order_by."),
    ("Detección y corrección N+1", "Prueba roja, prueba verde y reducción 23 a 2 con contenido idéntico."),
    ("Repositorio y observaciones", "Rama propia, preguntas vinculadas al código, costos y este documento."),
], [2.8, 3.7], size=8.8)

heading("Referencias", new_page=True)
for text in [
    "Guía de Laboratorio 07 ORM con IA duelo y verificador. Texto proporcionado por el estudiante, y rúbrica ORM en Django del campus.",
    "Django 6.0. Referencia de QuerySet, select_related, prefetch_related y relaciones. https://docs.djangoproject.com/en/6.0/ref/models/querysets/",
    "Django 6.0. Agregación y combinación de múltiples agregaciones. https://docs.djangoproject.com/en/6.0/topics/db/aggregation/",
    "Django 6.0. Herramientas de prueba y assertNumQueries. https://docs.djangoproject.com/en/6.0/topics/testing/tools/",
    "Python. HMAC y compare_digest. https://docs.python.org/3/library/hmac.html",
]:
    p(text)
heading("Evidencias para reproducir la revisión", 2)
table(["Archivo", "Qué permite comprobar"], [
    ("01-red-test.txt", "Fallo real inicial y las 23 sentencias capturadas."),
    ("02-green-test.txt", "La misma prueba después de la corrección."),
    ("03-full-suite.txt", "Resultado completo de las 31 pruebas."),
    ("dataset.json", "Conteos de la carga conocida."),
    ("duel-team.json y duel-ai.json", "Resultados, presupuestos, diferencias y SQL generado."),
    ("query-counts.json", "Comparación íntegra de los datos y consultas antes y después."),
    ("agent-results.json", "Planes, borrado aprobado, rechazo y restauración."),
    ("browser-results.json", "Comprobaciones de navegador, móvil y conservación de news."),
    ("Capturas PNG y JPG", "Panel, relaciones, comparación, N+1, acciones y editor."),
], [2.8, 3.7])
p("Todos los archivos anteriores se encuentran en deliverables/lab07/evidence dentro de la rama lab07-orm. Las capturas son de la ejecución local y los datos editoriales son ficticios de práctica.")

for element in (document._element, document.styles._element):
    for border in element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
target = OUTPUT / "Lab07_Diego_Magallanes_Linares.docx"
document.save(target)
print(target)
