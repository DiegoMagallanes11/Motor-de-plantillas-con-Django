"""Build the Spanish laboratory report from verified project evidence."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "deliverables"
EVIDENCE = OUTPUT / "evidence"
QA = OUTPUT / "qa"
QA.mkdir(parents=True, exist_ok=True)
document = Document()
section = document.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(.7)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.8)
section.right_margin = Inches(.8)
section.header_distance = Inches(.25)
section.footer_distance = Inches(.25)

for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
    style = document.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
normal = document.styles["Normal"]
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
document.styles["Title"].font.size = Pt(30)
document.styles["Heading 1"].font.size = Pt(20)
document.styles["Heading 2"].font.size = Pt(13)
document.styles["Heading 1"].paragraph_format.space_after = Pt(12)
document.styles["Heading 2"].paragraph_format.space_before = Pt(12)
document.core_properties.title = "Laboratorio 06 Plantillas y administrador de Django"
document.core_properties.author = "Diego Magallanes Linares"
document.core_properties.subject = "Desarrollo de Aplicaciones Empresariales"

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Laboratorio 06  |  ")
run.font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)


def p(text, style=None):
    return document.add_paragraph(text, style)


def heading(text, level=1, new_page=False):
    if new_page:
        document.add_page_break()
    document.add_heading(text, level)


def code(text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(text.strip("\n").splitlines()):
        run = paragraph.add_run(("\n" if index else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(8)
    return paragraph


def table(headers, rows, widths=None):
    result = document.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    if widths:
        for column, width in zip(result.columns, widths):
            column.width = Inches(width)
    for cell, text in zip(result.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(result.add_row().cells, row):
            cell.text = str(text)
    for index, row in enumerate(result.rows):
        for column, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "253E39" if index == 0 else ("F0F3F1" if index % 2 else "FFFFFF"))
            properties.append(shade)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "85")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            properties.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                border = OxmlElement(f"w:{edge}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "4")
                border.set(qn("w:color"), "D9D9D9")
                borders.append(border)
            properties.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    if index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return result


def picture(filename, caption, width=6.5, crop=None, max_height=5.2):
    source = EVIDENCE / filename
    if crop:
        image = Image.open(source)
        image = image.crop(crop)
        source = QA / f"crop-{Path(filename).stem}.png"
        image.save(source)
    image = Image.open(source)
    width = min(width, max_height * image.width / image.height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(source), width=Inches(width))
    paragraph = p(caption)
    paragraph.paragraph_format.space_after = Pt(9)
    for run in paragraph.runs:
        run.italic = True
        run.font.size = Pt(9)


p("GUÍA DE LABORATORIO 06", "Subtitle")
p("Plantillas y administrador de Django", "Title")
p("Portal de noticias Pulso", "Subtitle")
table(["Dato", "Información"], [
    ("Estudiante", "Diego Magallanes Linares"),
    ("Curso", "Desarrollo de Aplicaciones Empresariales"),
    ("Sección indicada en la guía", "4-C24-A"),
    ("Fecha de desarrollo", "15 de octubre de 2026"),
    ("Proyecto", "Pulso — Tecnología, Cultura y Ciencia"),
], [2.2, 4.3])
heading("Propósito y resultado", 2)
p("Se implementó un portal de noticias con una plantilla base, una tarjeta compartida y tres páginas conectadas: portada, detalle y listado por categoría. Los datos proceden de SQLite y pueden editarse desde Django Admin. Las diez pruebas automáticas y las comprobaciones del navegador finalizaron correctamente.")
heading("Capacidades", 2)
for text in [
    "Implementar el motor de plantillas con herencia y fragmentos reutilizables.",
    "Mostrar datos del modelo mediante variables, etiquetas de control y filtros.",
    "Gestionar contenido desde el administrador y publicarlo en las plantillas.",
]:
    p(text, "List Bullet")
heading("Participación y alcance", 2)
p("El trabajo corresponde a Diego Magallanes Linares y se desarrolló con asistencia de Codex para implementación, verificación y documentación. Se inició desde cero por no disponer del proyecto anterior. No se atribuye participación a otros integrantes. Las seis noticias, el autor editorial y las ilustraciones son material de ejemplo.")
p("El documento presenta los trece pasos de la guía, código relevante, evidencias reales, pruebas, observaciones y conclusiones. El código completo se encuentra en el repositorio indicado al final.")

heading("Preparación del proyecto", new_page=True)
heading("Paso 1 Entorno y aplicación news", 2)
p("Se utilizó Python 3.14.5 con Django 6.0.9 y Pillow 12.3.0. Django 6.0 admite Python 3.12, 3.13 y 3.14. El entorno virtual aísla las dependencias y requirements.txt fija sus versiones. La aplicación news concentra la responsabilidad editorial.")
code(r"""python -m venv .venv
.\.venv\Scripts\python.exe -m pip install -r requirements.txt
.\.venv\Scripts\python.exe scripts/setup_local.py
.\.venv\Scripts\python.exe manage.py runserver""")
p("Durante la preparación, el Python instalado no incorporó pip al entorno virtual. Se completó la instalación con python -m pip --python .\\.venv\\Scripts\\python.exe install -r requirements.txt. La incidencia no requirió modificar el funcionamiento del portal.")
heading("Paso 2 Plantillas estáticos y medios", 2)
p("config/settings.py registra news, busca las plantillas globales en templates y diferencia recursos estáticos de imágenes cargadas. La clave secreta se genera localmente y se lee desde el entorno; no se publica en el repositorio.")
code('''INSTALLED_APPS = [
    # Django built-in applications are also registered.
    "news.apps.NewsConfig",
]
TEMPLATES[0]["DIRS"] = [BASE_DIR / "templates"]
STATIC_URL = "static/"
STATICFILES_DIRS = [BASE_DIR / "static"]
STATIC_ROOT = BASE_DIR / "staticfiles"
MEDIA_URL = "media/"
MEDIA_ROOT = BASE_DIR / "media"
SECRET_KEY = os.environ["DJANGO_SECRET_KEY"]''')
p("El bloque anterior resume la configuración relevante; settings.py contiene también las aplicaciones, middleware y procesadores de contexto de Django. El archivo .env se excluye de Git y .env.example documenta las variables necesarias.")
heading("Recursos y normas aplicadas", 2)
p("Se trabajó en Windows con Python, Git y Visual Studio Code. Se utilizaron nombres de código en inglés y explicaciones en español, modelos en singular y migraciones versionadas. React no fue necesario para las páginas solicitadas. Se mantuvo el trabajo dentro del software y archivos del proyecto, sin intervenir hardware ni conexiones.")

heading("Datos y relaciones", new_page=True)
heading("Paso 3 Modelos migraciones y superusuario", 2)
p("Article representa una noticia; Author identifica a su autor y Category permite agrupar noticias. La relación de autor es uno a muchos y las categorías admiten una relación muchos a muchos. Se impide eliminar un autor que todavía tenga noticias asociadas.")
table(["Modelo", "Campos y relaciones principales"], [
    ("Author", "name, email y bio. Nombre visible y biografía editorial."),
    ("Category", "name, slug y description. Nombre único y ruta legible."),
    ("Article", "title, slug, summary, body, featured_image, published_at e is_published; author y categories."),
], [1.1, 5.4])
code('''class Article(models.Model):
    title = models.CharField("título", max_length=200)
    slug = models.SlugField(unique=True)
    summary = models.TextField("resumen")
    body = models.TextField("cuerpo")
    featured_image = models.ImageField(
        "imagen destacada", upload_to="articles/%Y/%m/", blank=True
    )
    published_at = models.DateTimeField(
        "fecha de publicación", default=timezone.now
    )
    is_published = models.BooleanField("publicado", default=True)
    author = models.ForeignKey(
        Author, on_delete=models.PROTECT, related_name="articles"
    )
    categories = models.ManyToManyField(
        Category, related_name="articles"
    )''')
p("La imagen es opcional: una tarjeta sin archivo utiliza un estado alternativo. Los slugs son únicos y evitan ambigüedades en las rutas. La fecha se almacena con zona horaria y el sitio utiliza America/Lima.")
code(r""".\.venv\Scripts\python.exe manage.py makemigrations news
.\.venv\Scripts\python.exe manage.py migrate
.\.venv\Scripts\python.exe scripts/setup_local.py""")
p("La migración news/migrations/0001_initial.py crea las tres tablas y sus relaciones. setup_local.py genera la configuración privada, aplica migraciones, carga datos y crea el superusuario diego con una contraseña aleatoria. La contraseña solo se conserva en .private/admin-access.txt, fuera de Git y de este documento.")

heading("Organización y herencia", new_page=True)
heading("Pasos 4 y 5 Base y fragmento compartido", 2)
p("base.html contiene la cabecera, navegación, pie y barra lateral. Sus bloques title, content y sidebar permiten cambiar cada página sin copiar esa estructura. El bloque heading adapta la presentación de portada y categoría.")
code('''{% load static %}
<title>{% block title %}Pulso{% endblock %}</title>
<link rel="stylesheet" href="{% static 'news/css/styles.css' %}">
<main id="content">{% block content %}{% endblock %}</main>
<aside>{% block sidebar %}...{% endblock %}</aside>''')
p("La tarjeta se define una sola vez en templates/news/_article_card.html. Tanto home.html como category.html la incluyen con el mismo contrato de contexto: article. El modificador only evita que dependa accidentalmente de otras variables de la página.")
code('''{% extends 'base.html' %}
{% block content %}
  {% for article in articles %}
    {% include 'news/_article_card.html' with article=article only %}
  {% empty %}
    <p>No hay noticias publicadas por el momento.</p>
  {% endfor %}
{% endblock %}''')
picture("10-editor.jpg", "Figura 1. Captura real de Visual Studio Code con la carpeta de plantillas y home.html.", max_height=3.9)
p("config reúne la configuración; news contiene datos, consultas y administrador; templates concentra la presentación y static la hoja de estilos. Esta separación permite localizar cada responsabilidad.")

heading("Portada con datos del modelo", new_page=True)
heading("Paso 6 Recorridos condiciones y filtros", 2)
p("La portada recibe el conjunto articles desde la vista. for genera una tarjeta por noticia; empty presenta un mensaje cuando no hay resultados. La tarjeta aplica date para la fecha y truncatewords:22 para limitar el resumen. Las seis publicaciones se distribuyen en tres categorías, con dos noticias en cada una.")
code('''{% if article.featured_image %}
  <img src="{{ article.featured_image.url }}" alt="">
{% else %}
  <div class="image-placeholder">pulso</div>
{% endif %}
<time>{{ article.published_at|date:"d M Y" }}</time>
<p>{{ article.summary|truncatewords:22 }}</p>''')
picture("01-home.png", "Figura 2. Portada real con datos de SQLite, imágenes cargadas y resúmenes recortados. Recorte superior de la captura completa.", crop=(0, 0, 1440, 1080), max_height=5)
p("La plantilla no decide qué registros son publicables. Esa consulta permanece en la vista y excluye borradores y fechas futuras. Los filtros de presentación se mantienen en la plantilla.")

heading("Detalle y listado por categoría", new_page=True)
heading("Pasos 7 y 8 Páginas especializadas", 2)
p("detail.html hereda la base y muestra título, resumen completo, imagen, autor, fecha, cuerpo y categorías. Los enlaces de categoría permiten continuar explorando el portal. category.html recibe la categoría seleccionada y reutiliza la misma tarjeta de la portada.")
code('''def category(request, slug):
    selected_category = get_object_or_404(Category, slug=slug)
    articles = published_articles().filter(categories=selected_category)
    return render(request, "news/category.html", {
        "category": selected_category, "articles": articles,
    })''')
picture("02-category.png", "Figura 3. Listado real de Tecnología con sus dos noticias. Ambas tarjetas usan el mismo fragmento de la portada.", max_height=4.5)
p("get_object_or_404 devuelve una respuesta 404 para un slug inexistente. published_articles utiliza select_related para el autor y prefetch_related para las categorías. La consulta filtra is_published=True y published_at menor o igual al momento actual.")

heading("Navegación y archivos", new_page=True)
heading("Paso 9 Rutas con nombres propios", 2)
p("news/urls.py define las tres rutas con el namespace news. Las plantillas usan la etiqueta url para generar enlaces, de modo que no dependen de direcciones internas escritas manualmente.")
code('''app_name = "news"
urlpatterns = [
    path("", views.home, name="home"),
    path("noticias/<slug:slug>/", views.detail, name="detail"),
    path("categorias/<slug:slug>/", views.category, name="category"),
]

<a href="{% url 'news:home' %}">Inicio</a>
<a href="{% url 'news:detail' article.slug %}">{{ article.title }}</a>
<a href="{% url 'news:category' item.slug %}">{{ item.name }}</a>''')
heading("Paso 10 CSS e imágenes", 2)
p("La base carga static y enlaza static/news/css/styles.css. Las imágenes generadas por Pillow se guardan en MEDIA_ROOT y se muestran con featured_image.url. config/urls.py habilita los medios solo cuando DEBUG está activo.")
code('''urlpatterns = [
    path("admin/", admin.site.urls),
    path("", include("news.urls")),
]
if settings.DEBUG:
    urlpatterns += static(
        settings.MEDIA_URL, document_root=settings.MEDIA_ROOT
    )''')
p("El navegador comprobó que las seis imágenes tenían dimensiones naturales mayores que cero y que el estilo de fondo coincidía con la hoja CSS. A un ancho de 390 píxeles no hubo desbordamiento horizontal. Las capturas completas, incluida la vista móvil, se conservan como evidencia en deliverables/evidence.")
heading("Caso sin resultados", 2)
p("Se verificó el caso vacío en las pruebas automáticas de portada y categoría. Además, se creó una categoría temporal sin noticias para capturar su mensaje en el navegador; después se retiró, manteniendo las tres categorías originales.")
if (EVIDENCE / "11-empty.png").exists():
    picture("11-empty.png", "Figura 4. Categoría temporal sin publicaciones y mensaje generado por empty.", crop=(80, 290, 1370, 800), max_height=2.3)

heading("Gestión editorial", new_page=True)
heading("Paso 11 Administrador de las tres entidades", 2)
p("Article, Category y Author se registran mediante admin.register. Cada entidad configura list_display, list_filter y search_fields. Article también utiliza campos de slug precompletados, selector horizontal de categorías y jerarquía de fechas.")
code('''@admin.register(Article)
class ArticleAdmin(admin.ModelAdmin):
    list_display = ("title", "author", "published_at", "is_published")
    list_filter = ("is_published", "categories", "author", "published_at")
    search_fields = ("title", "summary", "body")
    prepopulated_fields = {"slug": ("title",)}
    filter_horizontal = ("categories",)
    date_hierarchy = "published_at"''')
picture("06-admin-articles.png", "Figura 5. Listado real del administrador con seis noticias, búsqueda y filtros editoriales.", max_height=4.2)
p("La carga inicial se automatizó con seed_news para que otra instalación reproduzca las seis noticias y tres categorías. El panel permite crear nuevos registros, reemplazar imágenes y modificar todos esos datos. La búsqueda por el título de la noticia de accesibilidad devolvió un único resultado.")

heading("Cambios del panel al portal", new_page=True)
heading("Comprobación del paso 11", 2)
p("Se inició sesión en el administrador local, se abrió la noticia Diseñar herramientas que todos puedan usar y se añadió al resumen la frase «Actualización verificada desde el administrador.». Después de guardar, el navegador abrió el detalle y comprobó que el cambio era visible sin modificar el código.")
picture("08-admin-edit.png", "Figura 6. Formulario real de edición con campos de título, resumen y cuerpo. Recorte superior.", crop=(0, 0, 1440, 900), max_height=4.3)
picture("09-admin-change-visible.png", "Figura 7. El resumen publicado muestra la frase añadida desde el administrador. Recorte superior del detalle.", crop=(80, 150, 1050, 395), max_height=1.8)
p("La prueba finalizó restaurando el resumen original mediante el mismo formulario. El archivo browser-results.json registra adminEditVisible=true. También existe una prueba automática que envía el formulario administrativo y comprueba el título actualizado en la página pública.")

heading("Escapado automático", new_page=True)
heading("Paso 12 Guardar y mostrar una etiqueta HTML", 2)
p("En el cuerpo de la noticia de accesibilidad se guardó el siguiente texto. La base de datos conserva los caracteres de la etiqueta; la plantilla decide cómo representarlos en la respuesta HTML.")
code("<strong>Texto de prueba</strong>")
p("El detalle presenta literalmente <strong>Texto de prueba</strong>. El texto no se transforma en negrita y no se crea un elemento strong dentro del cuerpo de la noticia. El navegador verificó ambas condiciones.")
code('''<div class="article-body">{{ article.body|linebreaksbr }}</div>''')
picture("03-detail-escape.png", "Figura 8. El detalle conserva la etiqueta como texto visible. Recorte del cuerpo y de la biografía.", crop=(80, 825, 1050, 1330), max_height=3.8)
heading("Explicación del resultado", 2)
p("El escapado automático convierte caracteres especiales en entidades, como &lt; y &gt;, para que el navegador los muestre sin interpretar las etiquetas almacenadas. linebreaksbr conserva los saltos de línea y escapa el texto de entrada; no se utiliza safe ni se desactiva autoescape.")
p("Esto permite mostrar contenido redactado desde el administrador sin tratar automáticamente ese contenido como código HTML. El caso también está cubierto por test_detail_escapes_html_and_displays_relations, que comprueba la representación escapada y la ausencia del elemento HTML original.")

heading("Casos de prueba y resultados", new_page=True)
p("Las diez pruebas de Django se ejecutaron en una base temporal separada y finalizaron con OK. El chequeo del proyecto no informó errores y la revisión de migraciones no encontró cambios pendientes. Las comprobaciones del navegador se registraron por separado y no detectaron errores de página.")
table(["Caso y acción", "Resultado esperado y obtenido", "Evidencia"], [
    ("Abrir portada", "Seis tarjetas; esperado cumplido.", "01-home.png; navegador"),
    ("Seleccionar Tecnología", "Dos noticias relacionadas; cumplido.", "02-category.png; test categoría"),
    ("Abrir detalle", "Autor, fecha, categorías y cuerpo; cumplido.", "03-detail-escape.png; test detalle"),
    ("Listar sin noticias", "Mensaje de empty; cumplido en portada y categoría.", "Tests empty_home y empty_category"),
    ("Revisar resumen", "Recorte a 22 palabras; cumplido.", "test_home_uses_inheritance_card_and_filters"),
    ("Abrir CSS e imágenes", "Estilo aplicado y seis imágenes cargadas; cumplido.", "browser-results.json"),
    ("Vista móvil de 390 px", "Sin desbordamiento horizontal; cumplido.", "04-mobile.png; navegador"),
    ("Buscar y filtrar en Admin", "Solo registros coincidentes; cumplido.", "07-admin-search.png; test admin"),
    ("Editar desde Admin", "Cambio visible sin tocar código; cumplido.", "09-admin-change-visible.png"),
    ("Guardar HTML en cuerpo", "Etiqueta como texto; cumplido.", "Figura 8; test escapado"),
    ("Consultar borrador o futura", "No se publica y detalle devuelve 404; cumplido.", "test_drafts_future_and_missing_objects_return_404"),
    ("Abrir Admin sin sesión", "Redirección al acceso; cumplido.", "test_admin_requires_login"),
    ("Eliminar autor con noticias", "Se impide por PROTECT; cumplido.", "test_author_with_articles_is_protected"),
], [1.5, 3.0, 2.0])
code("Ran 10 tests in 0.695s\nOK\nSystem check identified no issues (0 silenced).")
p("Registro completo: deliverables/evidence/test-results.txt. Comprobaciones de navegador: deliverables/evidence/browser-results.json. Las capturas son resultados reales de la ejecución local; las noticias son datos ficticios de práctica.")

heading("Publicación y observaciones", new_page=True)
heading("Paso 13 Repositorio y entregable", 2)
p("El proyecto se publicó en la rama main del repositorio del estudiante. El primer commit de implementación es 5e662c1. El repositorio contiene modelos, vistas, plantillas, estilos, migraciones, dependencias, pruebas y README con instrucciones de instalación.")
p("Repositorio: https://github.com/DiegoMagallanes11/Motor-de-plantillas-con-Django")
p("El documento Word y las evidencias acompañan el proyecto. La subida de la guía al campus virtual corresponde al estudiante y permanece pendiente; no se presenta como una acción realizada.")
heading("Estructura de plantillas explicada", 2)
table(["Archivo", "Responsabilidad y reutilización"], [
    ("base.html", "Cabecera, navegación, pie y barra lateral; define bloques comunes."),
    ("news/_article_card.html", "Tarjeta única para portada y categoría; recibe article."),
    ("news/home.html", "Hereda base; recorre noticias y cubre portada vacía."),
    ("news/detail.html", "Hereda base; muestra una noticia y conserva el escapado."),
    ("news/category.html", "Hereda base; lista por categoría y reutiliza la tarjeta."),
], [2.15, 4.35])
heading("Observaciones de implementación", 2)
p("La barra lateral usa un procesador de contexto para consultar categorías sin repetir la consulta de navegación en cada vista. Los títulos y cuerpos editoriales proceden del modelo; los textos de interfaz comunes permanecen en las plantillas.")
p("Los archivos estáticos y medios tienen responsabilidades diferentes. La hoja CSS se versiona; las imágenes de práctica se reconstruyen con la carga inicial. La base SQLite, .env, contraseñas y entorno virtual se excluyen de Git.")
p("Se respetó la estructura de la sesión aunque el proyecto se inició desde cero. Para la captura del navegador se utilizó temporalmente el puerto 8003 y el runtime de trabajo por una incidencia de conexión del Python instalado. El procedimiento normal del README mantiene el servidor de desarrollo estándar.")

heading("Conclusiones", new_page=True)
p("1. La herencia permitió mantener una sola estructura común para las tres páginas. La tarjeta compartida demostró que portada y categoría pueden presentar noticias de manera consistente sin duplicar su marcado.")
p("2. Las variables, for, empty, if y los filtros date y truncatewords permitieron mostrar datos reales del modelo, formatear fechas y resolver listas vacías. La decisión de qué noticias publicar permaneció en las vistas, separada del HTML.")
p("3. El administrador permitió editar una noticia y observar el cambio en su detalle sin modificar código. La búsqueda y los filtros facilitaron localizar registros entre las seis noticias de ejemplo.")
p("4. La prueba con <strong>Texto de prueba</strong> confirmó que el escapado automático conserva la etiqueta como texto visible. Mantenerlo activo evita interpretar directamente como HTML el cuerpo almacenado.")
p("5. Las pruebas automáticas y las verificaciones del navegador aportaron evidencia de funcionamiento. El repositorio y las instrucciones reproducibles permiten reconstruir la base de datos y las imágenes en otra instalación sin compartir credenciales.")
heading("Correspondencia con la rúbrica", 2)
table(["Criterio de 5 puntos", "Evidencia disponible"], [
    ("Herencia y fragmentos", "base.html y tarjeta incluida en dos páginas; figuras 1 a 3."),
    ("Datos control y filtros", "Variables, for, empty, if, date y truncatewords; figuras 2 y 4 y pruebas."),
    ("Gestión desde Admin", "Tres entidades registradas y cambio visible; figuras 5 a 7."),
    ("Repositorio y observaciones", "Rama main, estructura explicada, README y este documento."),
], [2.5, 4])
heading("Referencias", 2)
for text in [
    "Guía de Laboratorio 06 Motor de plantillas con Django. Material proporcionado por el estudiante.",
    "Django 6.0 — Compatibilidad con Python. https://docs.djangoproject.com/en/6.0/releases/6.0/",
    "Django — Lenguaje de plantillas. https://docs.djangoproject.com/en/6.0/ref/templates/language/",
    "Django — Filtros incorporados. https://docs.djangoproject.com/en/6.0/ref/templates/builtins/",
    "Django — Administrador. https://docs.djangoproject.com/en/6.0/ref/contrib/admin/",
]:
    paragraph = p(text)
    for run in paragraph.runs:
        run.font.size = Pt(9)

target = OUTPUT / "Lab06_Diego_Magallanes_Linares.docx"
for element in (document._element, document.styles._element):
    for border in element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
document.save(target)
print(target)
